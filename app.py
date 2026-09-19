
import io
import os
import re
import hashlib
import tempfile
from pathlib import Path

import faiss
import gdown
import numpy as np
import streamlit as st
from docx import Document
from groq import Groq
from pypdf import PdfReader
from sentence_transformers import SentenceTransformer
from pptx import Presentation
from openpyxl import load_workbook
import xlrd

# ----------------------------
# App configuration
# ----------------------------
st.set_page_config(
    page_title="AI Document Assistant",
    page_icon="📄",
    layout="wide",
)

uploaded_files = st.file_uploader(
    "Upload PDF, Word, PowerPoint, Excel, TXT or MD files",
    type=["pdf", "docx", "txt", "md", "pptx", "xlsx", "xls", "csv"],
    accept_multiple_files=True,
)
EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
GROQ_MODEL = "openai/gpt-oss-20b"

# Common words that usually add little value to keyword search.
STOP_WORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "for", "from",
    "has", "have", "how", "i", "in", "is", "it", "of", "on", "or",
    "that", "the", "this", "to", "was", "were", "what", "when",
    "where", "which", "who", "why", "with", "you", "your"
}


# ----------------------------
# Cached model
# ----------------------------
@st.cache_resource
def load_embedding_model():
    """Load the Sentence Transformers model only once."""
    return SentenceTransformer(EMBEDDING_MODEL)


# ----------------------------
# Text extraction functions
# ----------------------------
def extract_pdf(file_bytes, filename):
    """Extract one text record per PDF page."""
    reader = PdfReader(io.BytesIO(file_bytes))
    records = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            records.append({
                "filename": filename,
                "page": page_number,
                "text": text.strip(),
            })

    return records


def extract_docx(file_bytes, filename):
    """Extract text from a DOCX file. Page numbers are not available reliably."""
    document = Document(io.BytesIO(file_bytes))
    text = "\n".join(
        paragraph.text for paragraph in document.paragraphs
        if paragraph.text.strip()
    )

    if not text.strip():
        return []

    return [{
        "filename": filename,
        "page": None,
        "text": text.strip(),
    }]


def extract_txt(file_bytes, filename):
    """Extract text from a TXT file."""
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="ignore")

    if not text.strip():
        return []

    return [{
        "filename": filename,
        "page": None,
        "text": text.strip(),
    }]


def extract_md(file_bytes, filename):
    """Extract Markdown as plain text."""
    return extract_txt(file_bytes, filename)
    
def extract_pptx(file_bytes, filename):
    presentation = Presentation(io.BytesIO(file_bytes))
    records = []

    for slide_no, slide in enumerate(presentation.slides, 1):
        text = "\n".join(
            shape.text for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        )

        if text.strip():
            records.append({
                "filename": filename,
                "page": slide_no,
                "text": text
            })

    return records


def extract_xlsx(file_bytes, filename):
    workbook = load_workbook(
        io.BytesIO(file_bytes),
        read_only=True,
        data_only=True
    )

    records = []

    for sheet in workbook.worksheets:
        rows = []

        for row in sheet.iter_rows(values_only=True):
            values = [str(v) for v in row if v is not None]
            if values:
                rows.append(" | ".join(values))

        if rows:
            records.append({
                "filename": filename,
                "page": None,
                "text": f"Sheet: {sheet.title}\n" + "\n".join(rows)
            })

    return records


def extract_xls(file_bytes, filename):
    workbook = xlrd.open_workbook(file_contents=file_bytes)
    records = []

    for sheet in workbook.sheets():
        rows = []

        for r in range(sheet.nrows):
            values = [
                str(sheet.cell_value(r, c))
                for c in range(sheet.ncols)
                if sheet.cell_value(r, c) != ""
            ]

            if values:
                rows.append(" | ".join(values))

        if rows:
            records.append({
                "filename": filename,
                "page": None,
                "text": f"Sheet: {sheet.name}\n" + "\n".join(rows)
            })

    return records


def extract_csv(file_bytes, filename):
    return extract_txt(file_bytes, filename)

def extract_file(file_bytes, filename):
    """Route a supported file to the correct extraction function."""
    extension = Path(filename).suffix.lower()

    if extension == ".pdf":
        return extract_pdf(file_bytes, filename)
    if extension == ".docx":
        return extract_docx(file_bytes, filename)
    if extension == ".txt":
        return extract_txt(file_bytes, filename)
    if extension == ".md":
        return extract_md(file_bytes, filename)
    if extension == ".pptx":
        return extract_pptx(file_bytes, filename)

    if extension == ".xlsx":
        return extract_xlsx(file_bytes, filename)

    if extension == ".xls":
        return extract_xls(file_bytes, filename)

    if extension == ".csv":
        return extract_csv(file_bytes, filename)

    raise ValueError(f"Unsupported file type: {extension}")


# ----------------------------
# Chunking
# ----------------------------
def split_text(text, chunk_size=1200, overlap=200):
    """Split text into overlapping character chunks."""
    text = re.sub(r"\s+", " ", text).strip()

    if not text:
        return []

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))

        # Try to end on a natural sentence/space boundary.
        if end < len(text):
            boundary = max(
                text.rfind(". ", start, end),
                text.rfind(" ", start, end)
            )
            if boundary > start + (chunk_size // 2):
                end = boundary + 1

        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = max(end - overlap, start + 1)

    return chunks


def create_chunks(extracted_records, chunk_size=1200, overlap=200):
    """Create chunks while preserving filename and page metadata."""
    chunks = []

    for record in extracted_records:
        text_chunks = split_text(
            record["text"],
            chunk_size=chunk_size,
            overlap=overlap,
        )

        for chunk_number, text_chunk in enumerate(text_chunks, start=1):
            chunks.append({
                "filename": record["filename"],
                "page": record["page"],
                "chunk_number": chunk_number,
                "text": text_chunk,
            })

    return chunks


# ----------------------------
# Embeddings and FAISS
# ----------------------------
def build_vector_store(chunks):
    """Create embeddings and a FAISS cosine-similarity index."""
    model = load_embedding_model()
    texts = [chunk["text"] for chunk in chunks]

    embeddings = model.encode(
        texts,
        convert_to_numpy=True,
        normalize_embeddings=True,
        show_progress_bar=False,
    ).astype("float32")

    index = faiss.IndexFlatIP(embeddings.shape[1])
    index.add(embeddings)

    return embeddings, index


# ----------------------------
# Search functions
# ----------------------------
def semantic_search(question, chunks, index, top_k=10):
    """Return top semantic matches from FAISS."""
    model = load_embedding_model()
    question_embedding = model.encode(
        [question],
        convert_to_numpy=True,
        normalize_embeddings=True,
        show_progress_bar=False,
    ).astype("float32")

    k = min(top_k, len(chunks))
    scores, indices = index.search(question_embedding, k)

    results = []
    for score, idx in zip(scores[0], indices[0]):
        if idx >= 0:
            results.append({
                "index": int(idx),
                "semantic_score": float(score),
            })

    return results


def important_words(text):
    """Extract simple meaningful keywords from a question."""
    words = re.findall(r"[A-Za-z0-9_'-]+", text.lower())
    return {
        word for word in words
        if len(word) >= 3 and word not in STOP_WORDS
    }


def keyword_search(question, chunks):
    """Score chunks using important word matches."""
    query_words = important_words(question)

    if not query_words:
        return []

    results = []

    for idx, chunk in enumerate(chunks):
        chunk_words = re.findall(
            r"[A-Za-z0-9_'-]+",
            chunk["text"].lower()
        )

        if not chunk_words:
            continue

        chunk_word_set = set(chunk_words)
        matched_words = query_words.intersection(chunk_word_set)

        if matched_words:
            coverage = len(matched_words) / len(query_words)
            frequency = sum(chunk_words.count(word) for word in matched_words)
            frequency_bonus = min(frequency / 10.0, 0.25)
            score = min(coverage + frequency_bonus, 1.0)

            results.append({
                "index": idx,
                "keyword_score": float(score),
            })

    return sorted(
        results,
        key=lambda item: item["keyword_score"],
        reverse=True,
    )


def hybrid_search(question, chunks, index, top_k=5):
    """
    Combine semantic and keyword search.

    Weighting:
    - 75% semantic similarity
    - 25% keyword match
    """
    semantic_results = semantic_search(
        question,
        chunks,
        index,
        top_k=min(20, len(chunks)),
    )
    keyword_results = keyword_search(question, chunks)

    combined = {}

    for item in semantic_results:
        idx = item["index"]
        # Cosine similarity is normally between -1 and 1.
        semantic_01 = (item["semantic_score"] + 1.0) / 2.0
        combined.setdefault(idx, {
            "semantic_score": 0.0,
            "keyword_score": 0.0,
        })
        combined[idx]["semantic_score"] = semantic_01

    for item in keyword_results:
        idx = item["index"]
        combined.setdefault(idx, {
            "semantic_score": 0.0,
            "keyword_score": 0.0,
        })
        combined[idx]["keyword_score"] = item["keyword_score"]

    ranked = []

    for idx, scores in combined.items():
        hybrid_score = (
            0.75 * scores["semantic_score"]
            + 0.25 * scores["keyword_score"]
        )

        ranked.append({
            **chunks[idx],
            "semantic_score": scores["semantic_score"],
            "keyword_score": scores["keyword_score"],
            "hybrid_score": hybrid_score,
        })

    ranked.sort(
        key=lambda item: item["hybrid_score"],
        reverse=True,
    )

    return ranked[:top_k]


# ----------------------------
# Groq answer generation
# ----------------------------
def answer_with_groq(question, retrieved_chunks):
    """Ask Groq to answer only from retrieved document context."""
    if "GROQ_API_KEY" not in st.secrets:
        raise RuntimeError(
            "GROQ_API_KEY is missing from Streamlit secrets."
        )

    client = Groq(api_key=st.secrets["GROQ_API_KEY"])

    context_parts = []
    for i, chunk in enumerate(retrieved_chunks, start=1):
        page_text = (
            f"Page {chunk['page']}"
            if chunk["page"] is not None
            else "Page unavailable"
        )

        context_parts.append(
            f"[Source {i}] {chunk['filename']} | {page_text}\n"
            f"{chunk['text']}"
        )

    context = "\n\n".join(context_parts)

    system_prompt = (
        "You are a document question-answering assistant. "
        "Answer ONLY from the supplied context. "
        "Do not use outside knowledge. "
        "If the context does not contain enough information, say: "
        "\"The information is not available in the provided documents.\" "
        "Keep the answer clear and concise."
    )

    user_prompt = (
        f"Question:\n{question}\n\n"
        f"Document context:\n{context}"
    )

    response = client.chat.completions.create(
        model=GROQ_MODEL,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        temperature=0.1,
        max_completion_tokens=1200,
    )

    return response.choices[0].message.content


# ----------------------------
# Google Drive loading
# ----------------------------
def is_drive_folder_link(url):
    return "/folders/" in url


def load_google_drive_files(url):
    """
    Download supported files from a public Google Drive file/folder link.

    The file or folder must be shared as "Anyone with the link".
    """
    collected_files = []

    with tempfile.TemporaryDirectory() as temp_dir:
        try:
            if is_drive_folder_link(url):
                downloaded_paths = gdown.download_folder(
                    url=url,
                    output=temp_dir,
                    quiet=True,
                    use_cookies=False,
                ) or []
            else:
                output_dir = temp_dir + os.sep
                downloaded_path = gdown.download(
                    url=url,
                    output=output_dir,
                    quiet=True,
                    use_cookies=False,
                )
                downloaded_paths = [downloaded_path] if downloaded_path else []
        except Exception as error:
            raise RuntimeError(
                "Could not download from Google Drive. "
                "Make sure the link is correct and shared as "
                "\"Anyone with the link\"."
            ) from error

        # For folder downloads, also scan the folder recursively in case
        # the downloader returns only some paths.
        all_paths = set()
        for path in downloaded_paths:
            if path:
                all_paths.add(str(path))

        for path in Path(temp_dir).rglob("*"):
            if path.is_file():
                all_paths.add(str(path))

        for path_string in sorted(all_paths):
            path = Path(path_string)
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                collected_files.append({
                    "name": path.name,
                    "bytes": path.read_bytes(),
                })

    return collected_files


# ----------------------------
# Processing helpers
# ----------------------------
def make_file_fingerprint(files):
    """Create a fingerprint so identical files are not reprocessed."""
    hasher = hashlib.sha256()

    for file_item in sorted(files, key=lambda item: item["name"]):
        hasher.update(file_item["name"].encode("utf-8"))
        hasher.update(file_item["bytes"])

    return hasher.hexdigest()


def process_files(files):
    """Run extraction -> chunking -> embeddings -> FAISS once."""
    extracted_records = []

    for file_item in files:
        records = extract_file(
            file_item["bytes"],
            file_item["name"],
        )
        extracted_records.extend(records)

    if not extracted_records:
        raise ValueError("No readable text was found in the documents.")

    chunks = create_chunks(extracted_records)

    if not chunks:
        raise ValueError("No text chunks could be created.")

    embeddings, index = build_vector_store(chunks)

    return {
        "files": files,
        "extracted_records": extracted_records,
        "chunks": chunks,
        "embeddings": embeddings,
        "index": index,
    }


def save_pipeline_if_new(files):
    """Process documents only when their fingerprint changes."""
    fingerprint = make_file_fingerprint(files)

    if (
        st.session_state.get("document_fingerprint") == fingerprint
        and st.session_state.get("pipeline") is not None
    ):
        return False

    pipeline = process_files(files)
    st.session_state.pipeline = pipeline
    st.session_state.document_fingerprint = fingerprint
    st.session_state.messages = []

    return True


# ----------------------------
# Session state
# ----------------------------
if "pipeline" not in st.session_state:
    st.session_state.pipeline = None

if "document_fingerprint" not in st.session_state:
    st.session_state.document_fingerprint = None

if "messages" not in st.session_state:
    st.session_state.messages = []


# ----------------------------
# User interface
# ----------------------------
st.title("📄 AI Document Assistant")
st.caption(
    "Upload documents or load public Google Drive files, then ask questions "
    "using hybrid semantic + keyword search."
)

with st.sidebar:
    st.header("1. Add documents")

    uploaded_files = st.file_uploader(
    "Upload PDF, Word, PowerPoint, Excel, TXT or MD files",
    type=["pdf", "docx", "txt", "md", "pptx", "xlsx", "xls", "csv"],
    accept_multiple_files=True,
)

    st.markdown("**OR**")

    drive_url = st.text_input(
        "Google Drive file or folder link",
        placeholder="https://drive.google.com/...",
    )

    process_button = st.button(
        "Process documents",
        type="primary",
        use_container_width=True,
    )

    if st.button("Clear documents", use_container_width=True):
        st.session_state.pipeline = None
        st.session_state.document_fingerprint = None
        st.session_state.messages = []
        st.rerun()


if process_button:
    files_to_process = []

    # Local uploads
    if uploaded_files:
        for uploaded_file in uploaded_files:
            files_to_process.append({
                "name": uploaded_file.name,
                "bytes": uploaded_file.getvalue(),
            })

    # Google Drive
    if drive_url.strip():
        with st.spinner("Loading files from Google Drive..."):
            try:
                drive_files = load_google_drive_files(drive_url.strip())
                files_to_process.extend(drive_files)
            except Exception as error:
                st.error(str(error))

    # Remove duplicate files with the same name + content.
    unique_files = {}
    for item in files_to_process:
        key = (
            item["name"],
            hashlib.sha256(item["bytes"]).hexdigest(),
        )
        unique_files[key] = item

    files_to_process = list(unique_files.values())

    if not files_to_process:
        st.warning("Please upload at least one document or add a Drive link.")
    else:
        with st.spinner("Extracting, chunking and embedding documents..."):
            try:
                was_new = save_pipeline_if_new(files_to_process)

                if was_new:
                    st.success("Documents processed successfully.")
                else:
                    st.info("These documents were already processed. Reusing them.")
            except Exception as error:
                st.error(f"Processing failed: {error}")


pipeline = st.session_state.pipeline

if pipeline:
    st.subheader("Document information")

    file_names = sorted({
        item["name"] for item in pipeline["files"]
    })

    pdf_page_counts = {}
    for record in pipeline["extracted_records"]:
        if record["page"] is not None:
            pdf_page_counts[record["filename"]] = max(
                pdf_page_counts.get(record["filename"], 0),
                record["page"],
            )

    info_rows = []
    for filename in file_names:
        file_records = [
            record for record in pipeline["extracted_records"]
            if record["filename"] == filename
        ]
        character_count = sum(
            len(record["text"]) for record in file_records
        )
        file_chunks = [
            chunk for chunk in pipeline["chunks"]
            if chunk["filename"] == filename
        ]

        info_rows.append({
            "Filename": filename,
            "Pages": pdf_page_counts.get(filename, "N/A"),
            "Characters": character_count,
            "Chunks": len(file_chunks),
        })

    st.dataframe(
        info_rows,
        use_container_width=True,
        hide_index=True,
    )

    metric1, metric2, metric3 = st.columns(3)
    metric1.metric("Documents", len(file_names))
    metric2.metric("Extracted sections", len(pipeline["extracted_records"]))
    metric3.metric("Created chunks", len(pipeline["chunks"]))

    with st.expander("Preview extracted text"):
        for record in pipeline["extracted_records"][:10]:
            label = record["filename"]
            if record["page"] is not None:
                label += f" — Page {record['page']}"

            st.markdown(f"**{label}**")
            st.text(record["text"][:1500])
            st.divider()

    st.subheader("Ask your documents")

    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

            if message.get("sources"):
                with st.expander("Retrieved sources"):
                    for i, source in enumerate(message["sources"], start=1):
                        page_label = (
                            f"Page {source['page']}"
                            if source["page"] is not None
                            else "Page unavailable"
                        )
                        st.markdown(
                            f"**Source {i}: {source['filename']} — {page_label}**"
                        )
                        st.caption(
                            f"Hybrid score: {source['hybrid_score']:.3f}"
                        )
                        st.write(source["text"])
                        st.divider()

    question = st.chat_input("Ask a question about your documents...")

    if question:
        st.session_state.messages.append({
            "role": "user",
            "content": question,
        })

        with st.chat_message("user"):
            st.markdown(question)

        with st.chat_message("assistant"):
            with st.spinner("Searching documents and generating answer..."):
                try:
                    retrieved = hybrid_search(
                        question,
                        pipeline["chunks"],
                        pipeline["index"],
                        top_k=5,
                    )

                    if not retrieved:
                        answer = (
                            "The information is not available in the provided documents."
                        )
                    else:
                        answer = answer_with_groq(question, retrieved)

                    st.markdown(answer)

                    with st.expander("Retrieved sources"):
                        for i, source in enumerate(retrieved, start=1):
                            page_label = (
                                f"Page {source['page']}"
                                if source["page"] is not None
                                else "Page unavailable"
                            )
                            st.markdown(
                                f"**Source {i}: {source['filename']} — {page_label}**"
                            )
                            st.caption(
                                f"Hybrid score: {source['hybrid_score']:.3f}"
                            )
                            st.write(source["text"])
                            st.divider()

                    st.session_state.messages.append({
                        "role": "assistant",
                        "content": answer,
                        "sources": retrieved,
                    })

                except Exception as error:
                    error_message = f"Answer generation failed: {error}"
                    st.error(error_message)

else:
    st.info(
        "Add documents in the sidebar and click **Process documents** "
        "to start."
    )
